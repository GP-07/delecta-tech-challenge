import copy
import json

import pytest
from openpyxl import load_workbook

from apps.menus import correction_services
from apps.menus.categorization_services import categorize_run_items
from apps.menus.models import FoodCategory, MenuItem, Method, StagedFoodCategory, StagedMenuItem
from apps.pipeline.steps import StepNotDoneError
from apps.venues import review_services
from apps.venues.models import StagedVenueMatch, VenueMatch
from libs.categorization.taxonomy import TaxonomyFileError
from tests.categorization.test_classifier import WordEmbedder
from tests.conftest import TIFLIS_ID
from tests.runs import matched_run, published_run, staged_run, succeeded

pytestmark = pytest.mark.django_db

TIFLIS_ITEMS = 89  # distinct items of its two menus (delivery and collection share them)
GEORGIAN_SALAD = "381eb832-ec40-4f0a-b40f-7a5d0b5c3d52"


def item(name, run=None):
    """The item staged by `run`, or the published one."""
    if run is not None:
        return StagedMenuItem.objects.select_related("category").get(run=run, name=name)
    return MenuItem.objects.select_related("category").get(name=name)


def test_every_item_of_a_linked_venue_is_staged_with_its_category(tiflis_expected):
    run = matched_run({TIFLIS_ID: tiflis_expected})

    result = categorize_run_items(run, embedder=None)

    assert (result.venues, result.items, result.created, result.classified, result.pending) == (
        1, TIFLIS_ITEMS, TIFLIS_ITEMS, TIFLIS_ITEMS, 0)
    assert StagedMenuItem.objects.filter(run=run).count() == TIFLIS_ITEMS
    assert not MenuItem.objects.exists()  # published only by publish_run
    salad = item("Ensalada Georgiana", run)
    assert (salad.category.name, salad.category.parent, salad.category.family) == ("salads", "salads", "food")
    assert (salad.method, salad.confidence) == (Method.NAME, 0.95)
    assert len(salad.menu_ids) == 2  # the same item in the delivery and the collection menus
    assert item("Patatas Fritas", run).category.name == "fries"
    assert item("Jagermeister, 1L", run).category.family == "drink"
    assert item("Khachapuri Adjaruli", run).category.name == "bread"  # from its section, "Panadería/Bakery"
    odjakhuri = item("Odjakhuri", run)
    assert (odjakhuri.category, odjakhuri.method) == (None, Method.UNKNOWN)


def test_the_run_stages_and_records_the_taxonomy_it_used(tiflis_expected):
    run = matched_run({TIFLIS_ID: tiflis_expected})

    categorize_run_items(run, embedder=None)

    run.refresh_from_db()
    assert len(run.food_categories_sha256) == 64
    assert StagedFoodCategory.objects.filter(run=run).count() == 101
    assert not FoodCategory.objects.exists()


def test_the_result_is_written_to_the_run_folder(tiflis_expected, tmp_path):
    run = matched_run({TIFLIS_ID: tiflis_expected}, crawl_dir=tmp_path / "run")

    result = categorize_run_items(run, embedder=None)

    assert result.output_file == str(tmp_path / "run" / "menu_items.json")
    rows = json.loads((tmp_path / "run" / "menu_items.json").read_text(encoding="utf-8"))
    assert len(rows) == TIFLIS_ITEMS
    salad = next(row for row in rows if row["item_name"] == "Ensalada Georgiana")
    assert salad | {"menu_ids": None, "item_id": None} == {
        "venue_id": TIFLIS_ID, "venue_name": "Tiflis Restaurant", "menu_ids": None, "item_id": None,
        "item_name": "Ensalada Georgiana", "section": "Ensaladas/Salads",
        "category_uidentifier": "f86dbf3dbbfe43508e3fb0c0a56e8d72a975be97", "name": "salads", "parent": "salads",
        "family": "food", "confidence": 0.95, "method": "name", "is_combo": False,
    }


def test_the_export_never_loads_the_venue_menus(tiflis_expected, django_assert_max_num_queries):
    from django.db import connection
    from django.test.utils import CaptureQueriesContext

    run = matched_run({TIFLIS_ID: tiflis_expected})
    with CaptureQueriesContext(connection) as queries:
        categorize_run_items(run, embedder=None)

    export = [query["sql"] for query in queries.captured_queries
              if '"menus_stagedmenuitem"."menu_ids"' in query["sql"] and "ORDER BY" in query["sql"]]
    assert export and all('"venues_stagedvenue"."menus"' not in sql for sql in export)


def test_categorization_needs_the_run_own_matching_to_have_succeeded(tiflis_expected):
    run = staged_run({TIFLIS_ID: tiflis_expected})

    with pytest.raises(StepNotDoneError, match="match_venues"):
        categorize_run_items(run, embedder=None)

    assert not StagedMenuItem.objects.exists()


def test_venues_in_review_are_left_out(tiflis_expected):
    run = matched_run({TIFLIS_ID: tiflis_expected})
    StagedVenueMatch.objects.update(status="REVIEW")

    result = categorize_run_items(run, embedder=None)

    assert (result.venues, result.items) == (0, 0)


def test_venues_confirmed_by_a_person_are_categorized(tiflis_expected):
    published_run({TIFLIS_ID: tiflis_expected})
    review_services.confirm(VenueMatch.objects.get(), user=None)
    run = matched_run({TIFLIS_ID: tiflis_expected})  # the venue is held: the run proposes no link for it

    assert not StagedVenueMatch.objects.filter(run=run).exists()
    assert categorize_run_items(run, embedder=None).items == TIFLIS_ITEMS


def test_only_venues_crawled_by_the_run_are_categorized(tiflis_expected):
    matched_run({TIFLIS_ID: tiflis_expected})

    result = categorize_run_items(matched_run({}), embedder=None)

    assert (result.venues, result.classified) == (0, 0)


def test_unchanged_items_keep_their_published_classification(tiflis_expected, settings):
    published_run({TIFLIS_ID: tiflis_expected})
    settings.CATEGORIZATION_MAX_ITEMS_PER_RUN = 1  # room for one: the others are not classified again
    run = matched_run({TIFLIS_ID: tiflis_expected})

    result = categorize_run_items(run, embedder=None)

    assert (result.created, result.removed, result.classified, result.pending) == (0, 0, 1, 0)
    published = dict(MenuItem.objects.values_list("item_id", "classified_at"))
    kept = [item_id for item_id, classified_at in StagedMenuItem.objects.filter(run=run).values_list(
        "item_id", "classified_at") if classified_at == published[item_id]]
    assert len(kept) == TIFLIS_ITEMS - 1


def test_a_category_corrected_by_a_person_is_kept_by_the_next_runs(tiflis_expected):
    published_run({TIFLIS_ID: tiflis_expected})
    odjakhuri = item("Odjakhuri")  # pork and potatoes, unknown to the rules
    correction_services.correct(odjakhuri, FoodCategory.objects.get(name="pork"), user=None)
    run = matched_run({TIFLIS_ID: tiflis_expected})

    result = categorize_run_items(run, embedder=None)

    staged = item("Odjakhuri", run)
    assert (staged.category.name, staged.method, staged.confidence) == ("pork", Method.MANUAL, 1.0)
    assert result.manual == 1


def test_items_gone_from_the_menu_are_left_out_and_changed_ones_classified_again(tiflis_expected):
    published_run({TIFLIS_ID: tiflis_expected})
    venue = copy.deepcopy(tiflis_expected)
    for menu in venue["menus"].values():
        for section in menu["sections"]:
            section["items"] = [item for item in section["items"] if item["name"] != "Patatas Fritas"]
            for entry in section["items"]:
                if entry["name"] == "Shoti":
                    entry["name"] = "Pan Shoti"
    run = matched_run({TIFLIS_ID: venue})

    result = categorize_run_items(run, embedder=None)

    assert (result.items, result.removed, result.created) == (TIFLIS_ITEMS - 1, 1, 0)
    assert not StagedMenuItem.objects.filter(run=run, name="Patatas Fritas").exists()
    assert item("Pan Shoti", run).category.name == "bread"


def test_items_over_the_run_limit_wait_for_the_next_run(tiflis_expected, settings):
    settings.CATEGORIZATION_MAX_ITEMS_PER_RUN = 50
    first = published_run({TIFLIS_ID: tiflis_expected})
    second = matched_run({TIFLIS_ID: tiflis_expected})

    result = categorize_run_items(second, embedder=None)

    assert MenuItem.objects.filter(run=first, method=Method.PENDING).count() == TIFLIS_ITEMS - 50
    assert (result.classified, result.pending) == (50, 0)  # the pending ones first


def taxonomy_without(category, settings, tmp_path):
    workbook = load_workbook(settings.FOOD_CATEGORIES_FILE)
    sheet = workbook.active
    sheet.delete_rows(next(row[0].row for row in sheet.iter_rows() if row[2].value == category))
    workbook.save(tmp_path / "food_categories.xlsx")
    return str(tmp_path / "food_categories.xlsx")


def test_a_category_removed_from_the_taxonomy_is_not_kept(tiflis_expected, settings, tmp_path):
    published_run({TIFLIS_ID: tiflis_expected})
    correction_services.correct(item("Plato de Queso"), FoodCategory.objects.get(name="other lacteos"), user=None)
    run = matched_run({TIFLIS_ID: tiflis_expected})
    run.food_categories_file = taxonomy_without("other lacteos", settings, tmp_path)
    run.save()

    categorize_run_items(run, embedder=None)

    cheese = item("Plato de Queso", run)
    assert (cheese.category.name, cheese.method) == ("cheese", Method.NAME)


def test_a_taxonomy_without_a_category_the_rules_use_fails_before_changing_anything(
        tiflis_expected, settings, tmp_path):
    run = matched_run({TIFLIS_ID: tiflis_expected})
    run.food_categories_file = taxonomy_without("salads", settings, tmp_path)
    run.save()

    with pytest.raises(TaxonomyFileError, match="salads"):
        categorize_run_items(run, embedder=None)

    assert not StagedMenuItem.objects.exists()


def test_a_run_without_taxonomy_file_fails(tiflis_expected):
    run = matched_run({TIFLIS_ID: tiflis_expected})
    run.food_categories_file = ""
    run.save()

    with pytest.raises(TaxonomyFileError):
        categorize_run_items(run, embedder=None)


def test_with_an_embedder_the_undecided_items_are_compared_with_the_decided_ones(tiflis_expected):
    result = categorize_run_items(matched_run({TIFLIS_ID: tiflis_expected}), embedder=WordEmbedder())

    assert result.embeddings is True
    assert result.items == TIFLIS_ITEMS


def test_a_step_retried_starts_again(tiflis_expected):
    run = matched_run({TIFLIS_ID: tiflis_expected})
    categorize_run_items(run, embedder=None)

    categorize_run_items(run, embedder=None)

    assert StagedMenuItem.objects.filter(run=run).count() == TIFLIS_ITEMS
    assert StagedFoodCategory.objects.filter(run=run).count() == 101
    succeeded(run, "categorize_menu_items")
