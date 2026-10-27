"""The food taxonomy (food_categories.xlsx): every category has a name, a parent and a family (food or drink).

The file is untrusted input: its size is capped, it is read in read-only mode (no formulas, no macros) and every row
is validated; a single invalid row rejects the whole file.
"""

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from pydantic import BaseModel, Field, ValidationError, field_validator

from libs.categorization.constants import GENERIC_PREFIX, TAXONOMY_HEADER, TAXONOMY_MAX_BYTES, TAXONOMY_MAX_ROWS


class TaxonomyFileError(Exception):
    """The taxonomy file is missing, too large or not a valid list of categories."""


class _Row(BaseModel):
    uidentifier: str = Field(pattern=r"^[0-9a-f]{40}$")
    family: str = Field(pattern=r"^(food|drink)$")
    name: str = Field(min_length=1, max_length=100)
    parent: str = Field(min_length=1, max_length=100)
    is_generic_or_other: bool

    @field_validator("name", "parent", "family", mode="before")
    @classmethod
    def _strip(cls, value):
        return value.strip().lower() if isinstance(value, str) else value

    @field_validator("is_generic_or_other", mode="before")
    @classmethod
    def _boolean(cls, value):
        return value.strip().lower() == "true" if isinstance(value, str) else value


@dataclass(frozen=True, slots=True)
class Category:
    uidentifier: str
    family: str
    name: str
    parent: str
    is_generic_or_other: bool

    @property
    def is_generic(self) -> bool:
        """'generic fish' (some fish), as opposed to 'other fish' (a fish not in the list)."""
        return self.is_generic_or_other and self.name.startswith(GENERIC_PREFIX)


class Taxonomy:
    def __init__(self, categories: list[Category]):
        self.categories = categories
        self._by_name = {category.name: category for category in categories}
        if len(self._by_name) != len(categories):
            raise TaxonomyFileError("category names must be unique")

    def __getitem__(self, name: str) -> Category:
        return self._by_name[name]

    def __contains__(self, name: str) -> bool:
        return name in self._by_name

    def __len__(self) -> int:
        return len(self.categories)

    def sibling(self, category: Category, prefix: str) -> Category | None:
        """'generic fish' -> 'other fish' (prefix 'other'), when the taxonomy has it."""
        if not category.is_generic_or_other:
            return None
        rest = category.name.split(" ", 1)[1] if " " in category.name else category.name
        return self._by_name.get(f"{prefix} {rest}")


def read_taxonomy(path: str | Path) -> Taxonomy:
    path = Path(path)
    if not path.is_file():
        raise TaxonomyFileError(f"food categories file not found: {path}")
    if path.stat().st_size > TAXONOMY_MAX_BYTES:
        raise TaxonomyFileError(f"food categories file larger than {TAXONOMY_MAX_BYTES} bytes")
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as error:  # openpyxl raises many types for a corrupt or non-xlsx file
        raise TaxonomyFileError(f"invalid food categories file: {error.__class__.__name__}") from error
    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        header = tuple(str(cell).strip().lower() for cell in next(rows, ()) if cell is not None)
        if header != TAXONOMY_HEADER:
            raise TaxonomyFileError(f"food categories file: expected columns {TAXONOMY_HEADER}, found {header}")
        categories = []
        for number, row in enumerate(rows, start=2):
            if number > TAXONOMY_MAX_ROWS + 1:
                raise TaxonomyFileError(f"food categories file has more than {TAXONOMY_MAX_ROWS} rows")
            if not any(row):
                continue
            try:
                categories.append(Category(**_Row(**dict(zip(TAXONOMY_HEADER, row, strict=False))).model_dump()))
            except ValidationError as error:
                raise TaxonomyFileError(f"food categories file, row {number}: {error.errors()[0]['msg']}") from error
    finally:
        workbook.close()
    if not categories:
        raise TaxonomyFileError("food categories file has no categories")
    return Taxonomy(categories)
